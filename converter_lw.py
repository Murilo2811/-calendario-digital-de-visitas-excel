"""Converte a PROGRAMAÇÃO L&W para o formato da base do app (Calendario_Digital_Base.xlsx).

Fonte: aba "Schedule 2025" (tabela estruturada; colunas batem 1:1 com Atividades).
Deriva Tecnicos e Clientes dos dados. Preserva Usuarios do arquivo base (para o login continuar valendo).

Uso:
    python converter_lw.py "PROGRAMAÇÃO L&W-2026 (1).xlsx" [Calendario_Digital_Base.xlsx] [saida.xlsx]

ponytail: só a aba tabular "Schedule 2025" é convertida. A grade "SCHEDULE BLOCKS"
não tem OS/HP/HT/HV e usa células mescladas -> conversão não-confiável, fica de fora.
Adicionar depois se precisar do 2026 vindo da grade.
"""
import sys
from datetime import datetime, date
import openpyxl

# Coluna (1-based) na aba "Schedule 2025" -> campo de Atividades. Header na linha 5, dados a partir da 6.
COLS = {
    "Semana": 1, "Cliente": 2, "Gerente": 3, "OS": 4, "Descricao": 5,
    "HP": 6, "HT": 7, "HV": 8, "Data Inicio": 9, "Data Fim": 10,
    "Tecnicos": 11, "Ultima Calibracao": 12, "Periodo": 13, "Status": 15,
}
ATIV_HEADER = ["ID", "Semana", "Cliente", "Gerente", "OS", "Descricao", "HP", "HT",
               "HV", "Data Inicio", "Data Fim", "Tecnicos", "Status",
               "Ultima Calibracao", "Periodo"]
DATA_START_ROW = 6


def num(v):
    """'28,5' / '5' -> float/int; vazio -> None."""
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return v
    s = str(v).strip().replace(",", ".")
    try:
        f = float(s)
        return int(f) if f.is_integer() else f
    except ValueError:
        return None


def d(v):
    """datetime -> date; passa date direto; resto vira None (ignora '***', '00:00:00' etc)."""
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    return None


def convert(src_path, base_path, out_path):
    wb = openpyxl.load_workbook(src_path, data_only=True)
    if "Schedule 2025" not in wb.sheetnames:
        raise SystemExit("Aba 'Schedule 2025' não encontrada em " + src_path)
    src = wb["Schedule 2025"]

    ativ, clientes, tecnicos = [], {}, {}
    sid = 0
    for row in src.iter_rows(min_row=DATA_START_ROW, values_only=True):
        get = lambda name: row[COLS[name] - 1] if COLS[name] - 1 < len(row) else None
        cliente = get("Cliente")
        ini, fim = d(get("Data Inicio")), d(get("Data Fim"))
        # linha só vale se tem cliente e pelo menos uma data (descarta placeholders/vazias)
        if not cliente or (ini is None and fim is None):
            continue
        cliente = str(cliente).strip()
        sid += 1
        tec = (str(get("Tecnicos")).strip() if get("Tecnicos") else "")
        ativ.append({
            "ID": f"s{sid}", "Semana": get("Semana"), "Cliente": cliente,
            "Gerente": get("Gerente"), "OS": get("OS"), "Descricao": get("Descricao"),
            "HP": num(get("HP")), "HT": num(get("HT")), "HV": num(get("HV")),
            "Data Inicio": ini, "Data Fim": fim, "Tecnicos": tec,
            "Status": get("Status"), "Ultima Calibracao": d(get("Ultima Calibracao")),
            "Periodo": num(get("Periodo")),
        })
        if cliente not in clientes:
            clientes[cliente] = f"c{len(clientes) + 1}"
        if tec and tec not in tecnicos:
            tecnicos[tec] = f"t{len(tecnicos) + 1}"

    out = openpyxl.Workbook()

    ws = out.active
    ws.title = "Atividades"
    ws.append(ATIV_HEADER)
    for a in ativ:
        ws.append([a[c] for c in ATIV_HEADER])

    ws = out.create_sheet("Tecnicos")
    ws.append(["ID", "Sigla", "Nome Completo", "Tipo", "Cor"])
    for sigla, tid in tecnicos.items():
        # ponytail: sem cadastro de nome/tipo na fonte -> default Internal. Ajustar no app.
        ws.append([tid, sigla, sigla, "Internal", "bg-red-100"])

    ws = out.create_sheet("Clientes")
    ws.append(["ID", "Nome", "Razao Social", "CNPJ", "Cidade", "Estado", "Contato", "Email", "Telefone"])
    for nome, cid in clientes.items():
        ws.append([cid, nome, "", "", "", "", "", "", ""])

    # Usuarios: preserva do base (hash de senha) se existir; senão cria aba vazia com header.
    ws = out.create_sheet("Usuarios")
    ws.append(["ID", "Usuario", "SenhaHash", "Papel", "NomeCompleto", "CriadoEm"])
    try:
        bwb = openpyxl.load_workbook(base_path, data_only=True)
        if "Usuarios" in bwb.sheetnames:
            for r in bwb["Usuarios"].iter_rows(min_row=2, values_only=True):
                if any(c is not None for c in r):
                    ws.append([d(c) if isinstance(c, datetime) else c for c in r])
    except FileNotFoundError:
        print(f"AVISO: base '{base_path}' não encontrada; aba Usuarios vazia.")

    out.save(out_path)
    print(f"OK -> {out_path}")
    print(f"  Atividades: {len(ativ)} | Clientes: {len(clientes)} | Tecnicos: {len(tecnicos)}")


if __name__ == "__main__":
    src = sys.argv[1] if len(sys.argv) > 1 else "PROGRAMAÇÃO L&W-2026 (1).xlsx"
    base = sys.argv[2] if len(sys.argv) > 2 else "Calendario_Digital_Base.xlsx"
    out = sys.argv[3] if len(sys.argv) > 3 else "Calendario_Digital_Convertido.xlsx"
    convert(src, base, out)
